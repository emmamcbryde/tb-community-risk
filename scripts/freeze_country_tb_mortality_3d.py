"""Read-only ABS workbook extraction to a NEW immutable validation snapshot.

No workbook is authored or modified. The100 terminal row remains source data;
it does not resolve the demographic100+ cell or authorize extrapolation.
"""
import argparse
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from engine.country_tb.data_common import canonical,digest,publish_immutable,seal,REPO_ROOT


def extract(path):
    from openpyxl import load_workbook
    b=Path(path).read_bytes();w=load_workbook(path,read_only=True,data_only=True);s=w['Table 9']
    assert 'Australia' in s['A3'].value and s['C6'].value=='qx' and s['G6'].value=='qx','ABS worksheet contract changed'
    source=[];derived=[]
    for row in s.iter_rows(min_row=8,max_row=108,max_col=9):
        age=row[0].value;assert type(age) is int
        for sex,offset in [('male',1),('female',5)]:
            raw={'age':age,'sex':sex,'lx':row[offset].value,'qx':str(row[offset+1].value),
                'Lx':row[offset+2].value,'ex':str(row[offset+3].value),
                'source_cell':'Table 9!'+row[offset+1].coordinate}
            source.append(raw);qx=Decimal(raw['qx']);assert 0<=qx<1
            if age<100:
                derived.append({'age_lower':age,'age_upper_exclusive':age+1,'sex':sex,'subgroup_id':None,
                    'survival_probability_one_year':str(1-qx),'source_cell':raw['source_cell']})
    w.close()
    assert len(source)==202 and len(derived)==200
    metadata={'source_id':'ABS_life_tables_2022_2024_Table9_Australia','geography_id':'who:AUS',
        'citation':'Australian Bureau of Statistics. Life expectancy2022–2024,Table1 data cube,Table9 Australia.',
        'url':'https://www.abs.gov.au/statistics/people/population/life-expectancy/2022-2024/3302055001DO001_20222024.xlsx',
        'source_bytes_sha256':hashlib.sha256(b).hexdigest(),'access_date':'2026-10-30','vintage':'2022–2024',
        'licence':{'name':'CC BY4.0','url':'https://www.abs.gov.au/website-privacy-copyright-and-disclaimer',
            'attribution':'Commonwealth of Australia,Australian Bureau of Statistics; selected mortality values and1-qx transformation',
            'review_status':'permission_checked_internal_validation_only'},
        'projection_status':'period_life_table_not_a_future_cohort_projection',
        'death_scope':'all_cause; removal_of_TB_deaths_for_cause_specific_competing_mortality_not_reviewed',
        'open_age_policy':'raw100 terminal row retained; no demographic100+ within-cell ages or beyond99 survival assumed',
        'adoption_status':'validation_snapshot_only_not_a_model_default_or_approved_competing_hazard'}
    value=seal({'schema_id':'country_tb.mortality_validation_snapshot.v1','metadata':metadata,'source_rows':source,'rows':derived,
        'production_available':False,'externally_validated':False})
    return value


def main():
    p=argparse.ArgumentParser();p.add_argument('--workbook',type=Path,required=True);p.add_argument('--output',type=Path,default=REPO_ROOT/'data/country_tb_validation/mortality');a=p.parse_args()
    value=extract(a.workbook);payload=canonical(value)+b'\n';identifier='abs-life2022-2024-'+value['profile_hash'][:12]
    manifest=seal({'schema_id':'country_tb.mortality_validation_manifest.v1','snapshot_id':identifier,
        'files':{'mortality.json':digest(payload)},'scientific_hash':value['profile_hash'],'immutable':True})
    target=publish_immutable(a.output,identifier,{'mortality.json':payload,'manifest.json':canonical(manifest)+b'\n'})
    print(json.dumps({'snapshot':str(target),'source_rows':202,'usable_one_year_intervals':200,'scientific_hash':value['profile_hash']}))


if __name__=='__main__':main()
