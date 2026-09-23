from __future__ import annotations

import importlib
from io import BytesIO
import sys
import time
import unittest
from pathlib import Path
from copy import deepcopy
from unittest.mock import patch

from openpyxl import load_workbook
import streamlit as st
from streamlit.testing.v1 import AppTest

from app.icon_arrays import build_100_person_visual_data, icon_grid_value
from app.results_page_display import (
    detailed_rows_for_display,
    format_interval_cells_for_display,
    key_metric_rows_for_display,
    results_rows_for_display,
)
from app.results_workbook import build_results_workbook, results_workbook_cache_key
from engine.apy.frozen_reference import load_frozen_reference_results


ROOT = Path(__file__).resolve().parents[1]


class ResultsPagePresentationTests(unittest.TestCase):
    def test_standard_results_page_omits_technical_information_sections(self) -> None:
        text = (ROOT / "pages" / "3_Results.py").read_text(encoding="utf-8")

        self.assertNotIn('expander("Technical information"', text)
        self.assertNotIn('expander("Additional technical information"', text)
        self.assertNotIn("Interface config", text)
        self.assertIn("build_results_workbook", text)

    def test_results_page_orders_key_metrics_before_details_and_per_100(self) -> None:
        text = (ROOT / "pages" / "3_Results.py").read_text(encoding="utf-8")

        self.assertLess(text.index('st.subheader("Key metrics")'), text.index('st.subheader("Detailed summary table")'))
        self.assertLess(text.index('st.subheader("Detailed summary table")'), text.index("What this means per 100 eligible people"))
        self.assertLess(text.index("What this means per 100 eligible people"), text.index('st.expander("Export results"'))

    def test_results_page_removes_redundant_sections(self) -> None:
        text = (ROOT / "pages" / "3_Results.py").read_text(encoding="utf-8")

        self.assertNotIn('st.subheader("Plain-language interpretation")', text)
        self.assertNotIn('st.subheader("Health Economics")', text)
        self.assertNotIn('st.subheader("Downloads")', text)
        self.assertIn('st.expander("Export results"', text)
        self.assertIn("Continue to Health Economics", text)

    def test_display_rows_use_friendly_interval_column_names(self) -> None:
        rows = results_rows_for_display(
            [
                {
                    "Metric": "nScreened",
                    "Median": 450.0,
                    "Low95": 450.0,
                    "High95": 450.0,
                }
            ]
        )

        self.assertEqual(
            list(rows[0].keys()),
            ["Outcome", "Median", "Low 95%", "High 95%"],
        )
        self.assertEqual(rows[0]["Outcome"], "People screened")
        self.assertEqual(rows[0]["Median"], 450.0)

    def test_deterministic_rows_use_expected_value_and_na_intervals(self) -> None:
        rows = results_rows_for_display(
            [
                {
                    "Metric": "nScreened",
                    "Median": 450.0,
                    "Low95": 450.0,
                    "High95": 450.0,
                }
            ],
            model_type="expected_value",
        )

        self.assertEqual(
            list(rows[0].keys()),
            ["Outcome", "Expected value", "Low 95%", "High 95%"],
        )
        self.assertEqual(rows[0]["Expected value"], 450.0)
        self.assertIsNone(rows[0]["Low 95%"])
        self.assertIsNone(rows[0]["High 95%"])
        formatted = format_interval_cells_for_display(rows)
        self.assertEqual(formatted[0]["Low 95%"], "N/A")
        self.assertEqual(formatted[0]["High 95%"], "N/A")

    def test_rendered_deterministic_results_page_shows_na_intervals(self) -> None:
        if not hasattr(st, "secrets"):
            importlib.reload(st)
        sys.modules["streamlit"] = st
        app = AppTest.from_file(str(ROOT / "pages" / "3_Results.py"))
        app.session_state["results_bundle"] = {
            "metadata": {
                "modelType": "expected_value",
                "analysisBasis": "sa_health_matlab_v9_compatibility_reference",
                "naturalHistorySemantics": "matlab_v9_implicit_early_late",
                "scenarioLabel": "Deterministic check",
            },
            "headline": {
                "keyMetricsRows": [
                    {"Metric": "nScreened", "Median": 450.0, "Low95": 450.0, "High95": 450.0}
                ],
                "summaryRows": [
                    {"Metric": "nScreened", "Median": 450.0, "Low95": 450.0, "High95": 450.0}
                ],
            },
            "technical": {
                "eventLedger": {
                    "metadata": {
                        "modelType": "expected_value",
                        "analysisBasis": "sa_health_matlab_v9_compatibility_reference",
                        "naturalHistorySemantics": "matlab_v9_implicit_early_late",
                    }
                },
                "interfaceConfig": {},
            },
            "downloads": {},
        }
        app.session_state["economics_config"] = {}

        app.run(timeout=30)

        self.assertFalse(app.exception)
        self.assertIn(
            "Simulation intervals are not applicable to a single deterministic run.",
            [caption.value for caption in app.caption],
        )
        rendered = app.dataframe[0].value
        self.assertEqual(list(rendered.columns), ["Outcome", "Expected value", "Low 95%", "High 95%"])
        self.assertEqual(rendered.loc[0, "Low 95%"], "N/A")
        self.assertEqual(rendered.loc[0, "High 95%"], "N/A")

    def test_results_page_does_not_build_workbook_during_render(self) -> None:
        start = time.perf_counter()
        with patch(
            "app.results_workbook.build_results_workbook",
            side_effect=AssertionError("Workbook should not be built during ordinary Results rendering"),
        ):
            app = AppTest.from_file(str(ROOT / "pages" / "3_Results.py"), default_timeout=120)
            app.run(timeout=120)
        elapsed = time.perf_counter() - start

        self.assertFalse(app.exception)
        self.assertLess(elapsed, 30.0)
        self.assertIn("Prepare Excel workbook", [button.label for button in app.button])
        self.assertNotIn("prepared_results_workbook", app.session_state)

    def test_results_page_prepares_parseable_workbook_only_after_click_and_reuses_cache(self) -> None:
        with patch("app.results_workbook.build_results_workbook", wraps=build_results_workbook) as wrapped:
            app = AppTest.from_file(str(ROOT / "pages" / "3_Results.py"), default_timeout=120)
            app.run(timeout=120)
            self.assertEqual(wrapped.call_count, 0)

            next(button for button in app.button if button.label == "Prepare Excel workbook").click().run(timeout=120)
            self.assertEqual(wrapped.call_count, 1)
            prepared = app.session_state["prepared_results_workbook"]
            payload = prepared["bytes"]
            self.assertGreater(len(payload), 1000)

            workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            try:
                self.assertIn("Headline_results", workbook.sheetnames)
                self.assertIn("Economic_annual_by_arm", workbook.sheetnames)
                self.assertIn("Economic_replicates", workbook.sheetnames)
                headline_rows = list(workbook["Headline_results"].iter_rows(values_only=True))
                self.assertGreater(len(headline_rows), 1)
                annual_rows = list(workbook["Economic_annual_by_arm"].iter_rows(values_only=True))
                annual_headers = set(annual_rows[0])
                self.assertIn("Workbook export note", annual_headers)
                self.assertLess(len(annual_rows), 500)
            finally:
                workbook.close()

            app.run(timeout=120)
            self.assertEqual(wrapped.call_count, 1)
            self.assertEqual(app.session_state["prepared_results_workbook"]["bytes"], payload)
            self.assertTrue(
                any(
                    getattr(item, "type", None) == "download_button"
                    and getattr(item, "label", "") == "Download consolidated results workbook"
                    for item in app
                )
            )

    def test_results_workbook_cache_key_changes_with_economic_and_epidemiological_inputs(self) -> None:
        payload = load_frozen_reference_results()
        bundle = payload["resultsBundle"]
        economics = payload["referenceEconomics"]
        economics_config = payload["economicsConfig"]

        base_key = results_workbook_cache_key(
            bundle=bundle,
            economics_results=economics,
            economics_config=economics_config,
            results_stale=False,
            dirty_economics=False,
            decision_analysis_results={},
        )
        changed_economics = deepcopy(economics_config)
        changed_economics.setdefault("metadata", {})["unit-test-change"] = "changed"
        changed_bundle = deepcopy(bundle)
        changed_bundle.setdefault("metadata", {})["configurationHash"] = "different-configuration"
        deterministic_bundle = deepcopy(bundle)
        deterministic_bundle.setdefault("metadata", {})["modelType"] = "expected_value"
        deterministic_bundle["metadata"]["analysisMethod"] = "expected_value"

        self.assertNotEqual(
            base_key,
            results_workbook_cache_key(
                bundle=bundle,
                economics_results=economics,
                economics_config=changed_economics,
                results_stale=False,
                dirty_economics=False,
                decision_analysis_results={},
            ),
        )
        self.assertNotEqual(
            base_key,
            results_workbook_cache_key(
                bundle=changed_bundle,
                economics_results=economics,
                economics_config=economics_config,
                results_stale=False,
                dirty_economics=False,
                decision_analysis_results={},
            ),
        )
        self.assertNotEqual(
            base_key,
            results_workbook_cache_key(
                bundle=deterministic_bundle,
                economics_results=economics,
                economics_config=economics_config,
                results_stale=False,
                dirty_economics=False,
                decision_analysis_results={},
            ),
        )

    def test_results_page_clears_prepared_workbook_for_stale_results(self) -> None:
        payload = load_frozen_reference_results()
        app = AppTest.from_file(str(ROOT / "pages" / "3_Results.py"), default_timeout=120)
        app.session_state["results_bundle"] = payload["resultsBundle"]
        app.session_state["economics_results"] = payload["referenceEconomics"]
        app.session_state["economics_config"] = payload["economicsConfig"]
        app.session_state["prepared_results_workbook"] = {"cacheKey": "old", "bytes": b"old", "fileName": "old.xlsx"}
        app.session_state["results_stale"] = True

        app.run(timeout=120)

        self.assertFalse(app.exception)
        self.assertNotIn("prepared_results_workbook", app.session_state)

    def test_key_metrics_include_active_tb_comparator_intervention_and_reduction(self) -> None:
        key_rows = [{"Metric": "nScreened", "Median": 450, "Low95": 450, "High95": 450}]
        dynamic_rows = [
            {"Metric": "cumulative_baseline_active_tb_cases", "Median": 37, "Low95": 26, "High95": 50},
            {"Metric": "cumulative_intervention_active_tb_cases", "Median": 25, "Low95": 15, "High95": 38},
            {"Metric": "cumulative_cases_averted", "Median": 12, "Low95": 4, "High95": 22},
            {"Metric": "relative_reduction_cumulative_active_tb_cases", "Median": 0.32, "Low95": 0.12, "High95": 0.48},
        ]

        labels = {
            row["Outcome"]
            for row in key_metric_rows_for_display(key_rows, dynamic_rows, model_type="agent_based")
        }

        self.assertIn("People screened", labels)
        self.assertIn("Comparator active TB", labels)
        self.assertIn("Intervention active TB", labels)
        self.assertIn("Active TB cases averted", labels)
        self.assertIn("Relative reduction in active TB", labels)

    def test_active_tb_averted_key_metric_falls_back_without_dynamic_rows(self) -> None:
        rows = key_metric_rows_for_display(
            [
                {"Metric": "nScreened", "Median": 450, "Low95": 450, "High95": 450},
                {"Metric": "nPreventedActiveTB", "Median": 12, "Low95": 4, "High95": 22},
            ],
            [],
            model_type="agent_based",
        )

        labels = [row["Outcome"] for row in rows]
        self.assertIn("Active TB cases averted", labels)
        self.assertEqual(labels.count("Active TB cases averted"), 1)

    def test_detailed_summary_adds_non_key_rows(self) -> None:
        key_rows = [{"Metric": "nScreened", "Median": 450, "Low95": 450, "High95": 450}]
        summary_rows = [
            *key_rows,
            {"Metric": "nADRstop", "Median": 3.0, "Low95": 0.0, "High95": 7.0},
        ]

        details = detailed_rows_for_display(summary_rows, key_rows, model_type="agent_based")

        self.assertEqual(len(details), 1)
        self.assertEqual(details[0]["Outcome"], "ADR-related stops")

    def test_per_100_uses_unrounded_underlying_values_and_explicit_denominator(self) -> None:
        rows = build_100_person_visual_data(
            {
                "replicateTotals": [
                    {"arm": "comparator", "eventName": "eligible_population", "value": 1500},
                    {"arm": "intervention", "eventName": "eligible_population", "value": 1500},
                    {"arm": "intervention", "eventName": "active_tb_cases_prevented", "value": 11.9665},
                ]
            },
            outcomes=["active_tb_cases_prevented"],
        )

        self.assertEqual(
            rows[0]["intervention"]["displayPer100"],
            "0.8 per 100",
        )
        self.assertAlmostEqual(rows[0]["intervention"]["exactPer100"], 11.9665 / 1500 * 100)

    def test_per_100_scales_when_population_size_changes(self) -> None:
        rows = build_100_person_visual_data(
            {
                "replicateTotals": [
                    {"arm": "comparator", "eventName": "eligible_population", "value": 3000},
                    {"arm": "intervention", "eventName": "eligible_population", "value": 3000},
                    {"arm": "intervention", "eventName": "active_tb_cases_prevented", "value": 12},
                ]
            },
            outcomes=["active_tb_cases_prevented"],
        )

        self.assertEqual(rows[0]["intervention"]["displayPer100"], "0.4 per 100")

    def test_very_small_nonzero_per_100_value_keeps_fractional_signal(self) -> None:
        value = icon_grid_value(0.04)

        self.assertEqual(value["displayPer100"], "0.04 per 100")
        self.assertEqual(value["filledIcons"], 0)


if __name__ == "__main__":
    unittest.main()
